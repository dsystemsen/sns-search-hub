"""Excel report builder.

Generates a multi-sheet workbook for a channel over a date range:
- サマリー: KPI totals + subscriber net
- 日次推移: per-day metrics
- 動画別ランキング: top videos by views
"""

from __future__ import annotations

import io
from datetime import date
from decimal import Decimal

from django.db.models import Sum
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from apps.channels.models import Channel

from .models import ChannelDailyMetric, VideoDailyMetric


HEADER_FILL = PatternFill("solid", fgColor="1F2937")
HEADER_FONT = Font(bold=True, color="FFFFFF")
TITLE_FONT = Font(bold=True, size=14)


def _autosize(ws):
    for col in ws.columns:
        max_len = 0
        column = col[0].column_letter
        for cell in col:
            value = "" if cell.value is None else str(cell.value)
            max_len = max(max_len, len(value))
        ws.column_dimensions[column].width = min(max(max_len + 2, 10), 40)


def _style_header(ws, row: int, last_col: int):
    for col in range(1, last_col + 1):
        cell = ws.cell(row=row, column=col)
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL
        cell.alignment = Alignment(horizontal="center", vertical="center")


def build_channel_report(channel: Channel, start: date, end: date) -> bytes:
    wb = Workbook()

    # === Summary sheet ===
    summary = wb.active
    summary.title = "サマリー"
    summary["A1"] = f"{channel.name} レポート"
    summary["A1"].font = TITLE_FONT
    summary["A2"] = f"期間: {start.isoformat()} 〜 {end.isoformat()}"

    totals = VideoDailyMetric.objects.filter(
        video__channel=channel, date__range=(start, end)
    ).aggregate(
        views=Sum("views"),
        watch_time_minutes=Sum("watch_time_minutes"),
        likes=Sum("likes"),
        comments=Sum("comments"),
        shares=Sum("shares"),
        estimated_revenue_usd=Sum("estimated_revenue_usd"),
    )

    channel_metrics = ChannelDailyMetric.objects.filter(
        channel=channel, date__range=(start, end)
    ).order_by("date")
    subs_gained = sum(m.subscribers_gained for m in channel_metrics)
    subs_lost = sum(m.subscribers_lost for m in channel_metrics)
    latest = channel_metrics.last()

    summary_rows = [
        ("KPI", "値"),
        ("再生回数", totals["views"] or 0),
        ("視聴時間 (分)", float(totals["watch_time_minutes"] or 0)),
        ("高評価", totals["likes"] or 0),
        ("コメント", totals["comments"] or 0),
        ("共有", totals["shares"] or 0),
        ("推定収益 (USD)", float(totals["estimated_revenue_usd"] or 0)),
        ("登録者総数", latest.subscribers_total if latest else "—"),
        ("登録者増加", subs_gained),
        ("登録者減少", subs_lost),
        ("登録者純増", subs_gained - subs_lost),
    ]
    start_row = 4
    for i, (k, v) in enumerate(summary_rows):
        summary.cell(row=start_row + i, column=1, value=k)
        summary.cell(row=start_row + i, column=2, value=v)
    _style_header(summary, start_row, 2)
    _autosize(summary)

    # === Daily sheet ===
    daily = wb.create_sheet("日次推移")
    daily_headers = ["日付", "再生回数", "視聴時間(分)", "高評価", "コメント", "共有", "登録者総数"]
    for col, h in enumerate(daily_headers, start=1):
        daily.cell(row=1, column=col, value=h)
    _style_header(daily, 1, len(daily_headers))

    daily_qs = (
        VideoDailyMetric.objects.filter(video__channel=channel, date__range=(start, end))
        .values("date")
        .annotate(
            views=Sum("views"),
            watch_time_minutes=Sum("watch_time_minutes"),
            likes=Sum("likes"),
            comments=Sum("comments"),
            shares=Sum("shares"),
        )
        .order_by("date")
    )
    channel_daily = {m.date: m for m in channel_metrics}
    for i, d in enumerate(daily_qs, start=2):
        cm = channel_daily.get(d["date"])
        daily.cell(row=i, column=1, value=d["date"].isoformat())
        daily.cell(row=i, column=2, value=d["views"] or 0)
        daily.cell(row=i, column=3, value=float(d["watch_time_minutes"] or 0))
        daily.cell(row=i, column=4, value=d["likes"] or 0)
        daily.cell(row=i, column=5, value=d["comments"] or 0)
        daily.cell(row=i, column=6, value=d["shares"] or 0)
        daily.cell(row=i, column=7, value=cm.subscribers_total if cm else None)
    _autosize(daily)

    # === Top videos sheet ===
    top = wb.create_sheet("動画別ランキング")
    top_headers = ["順位", "動画タイトル", "Video ID", "再生回数", "視聴時間(分)", "高評価", "コメント"]
    for col, h in enumerate(top_headers, start=1):
        top.cell(row=1, column=col, value=h)
    _style_header(top, 1, len(top_headers))

    top_qs = (
        VideoDailyMetric.objects.filter(video__channel=channel, date__range=(start, end))
        .values("video__title", "video__youtube_video_id")
        .annotate(
            views=Sum("views"),
            watch_time_minutes=Sum("watch_time_minutes"),
            likes=Sum("likes"),
            comments=Sum("comments"),
        )
        .order_by("-views")[:50]
    )
    for i, v in enumerate(top_qs, start=2):
        top.cell(row=i, column=1, value=i - 1)
        top.cell(row=i, column=2, value=v["video__title"])
        top.cell(row=i, column=3, value=v["video__youtube_video_id"])
        top.cell(row=i, column=4, value=v["views"] or 0)
        top.cell(row=i, column=5, value=float(v["watch_time_minutes"] or 0))
        top.cell(row=i, column=6, value=v["likes"] or 0)
        top.cell(row=i, column=7, value=v["comments"] or 0)
    _autosize(top)

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
